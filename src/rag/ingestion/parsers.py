from dataclasses import dataclass
from typing import Any

from rag.core.config import get_settings
from rag.ingestion import sniff


@dataclass
class Page:
    number: int | None
    markdown: str
    ocr: bool = False  # text came from OCR of an image or scanned page


def _pdf(data: bytes) -> list[Page]:
    import pymupdf
    import pymupdf4llm

    with pymupdf.open(stream=data, filetype="pdf") as doc:  # type: ignore[no-untyped-call]
        if doc.needs_pass:
            raise ValueError("Encrypted PDFs are not supported")
        if len(doc) > get_settings().max_pdf_pages:
            raise ValueError("PDF exceeds the allowed page count")
        if not any(page.get_text().strip() for page in doc):
            return _scanned_pdf(doc)
        pages = pymupdf4llm.to_markdown(doc, page_chunks=True, show_progress=False)
    return [Page(number=int(p["metadata"].get("page", i + 1)), markdown=p["text"]) for i, p in enumerate(pages)]


def _scanned_pdf(doc: Any) -> list[Page]:
    from rag.receipts.ocr import ocr, render_pdf_page

    pages = (doc[i] for i in range(min(len(doc), get_settings().max_ocr_pages)))
    return [Page(number=i + 1, markdown=ocr(render_pdf_page(page)), ocr=True) for i, page in enumerate(pages)]


def _image(data: bytes) -> list[Page]:
    from rag.receipts.ocr import load_image, ocr

    return [Page(number=1, markdown=ocr(load_image(data)), ocr=True)]


def _docx(data: bytes) -> list[Page]:
    import io

    import docx

    document = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name if para.style is not None else "") or ""
        if style.startswith("Heading ") and style[8:].isdigit():
            lines.append(f"{'#' * min(int(style[8:]), 6)} {text}")
        elif style == "Title":
            lines.append(f"# {text}")
        else:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return [Page(number=None, markdown="\n\n".join(lines))]


def _html(data: bytes) -> list[Page]:
    import trafilatura

    # Parses the uploaded bytes only; never follows links or fetches remote resources (SSRF-safe).
    text = trafilatura.extract(
        data.decode("utf-8", errors="replace"), output_format="markdown", include_links=False, include_images=False
    )
    return [Page(number=None, markdown=text or "")]


def parse(data: bytes, mime: str) -> list[Page]:
    if mime == sniff.PDF:
        return _pdf(data)
    if mime == sniff.DOCX:
        return _docx(data)
    if mime == sniff.HTML:
        return _html(data)
    if mime in sniff.IMAGES:
        return _image(data)
    return [Page(number=None, markdown=data.decode("utf-8", errors="replace"))]
